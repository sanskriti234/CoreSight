# =========================================================
#              CoreSight - Management Module
# =========================================================

import os
import json
import shutil
import sqlite3
from datetime import datetime
from fastapi import APIRouter, Form
from fastapi.responses import JSONResponse
from face_service import known_encodings, known_rolls

# =========================================================
# Router
# =========================================================
router = APIRouter(prefix="", tags=["Management"])

# =========================================================
# Shared Config (same as main.py)
# =========================================================
<<<<<<< HEAD
BASE_DIR = r"D:\CoreSight"
from paths import ATTENDANCE_DIR, ENCODINGS_DIR, IMAGES_DIR, DETAILS_DIR, DOCUMENTS_DIR
=======
from paths import BASE_DIR,ATTENDANCE_DIR, ENCODINGS_DIR, IMAGES_DIR, DETAILS_DIR, DOCUMENTS_DIR,FRONTEND_DIR
>>>>>>> ea91db6 (Face recogntion and document module integrated)


DB_PATH = "Core.db"
conn = sqlite3.connect(DB_PATH, check_same_thread=False)
cursor = conn.cursor()

# =========================================================
# Helper: Reload Encodings
# =========================================================
def reload_encodings(known_encodings: list, known_rolls: list):
    known_encodings.clear()
    known_rolls.clear()

    if os.path.exists(ENCODINGS_DIR):
        for file in os.listdir(ENCODINGS_DIR):
            if file.endswith(".pkl"):
                import pickle, numpy as np
                with open(os.path.join(ENCODINGS_DIR, file), "rb") as f:
                    data = pickle.load(f)

                if isinstance(data, dict):
                    known_encodings.append(data["encoding"])
                    known_rolls.append(data["roll_no"])
                else:
                    known_encodings.append(np.array(data))
                    known_rolls.append(os.path.splitext(file)[0])

# =========================================================
# GET: All Students
# =========================================================
@router.get("/students")
def get_all_students():
    cursor.execute("""
        SELECT student_name, roll_no, email, phone, department, academic_year
        FROM students
        ORDER BY roll_no
    """)
    rows = cursor.fetchall()

    return {
        "status": "success",
        "count": len(rows),
        "students": [
            {
                "name": r[0],
                "roll_no": r[1],
                "email": r[2],
                "phone": r[3],
                "department": r[4],
                "academic_year": r[5]
            } for r in rows
        ]
    }

# =========================================================
# POST: Update Student
# =========================================================
@router.post("/update-student")
def update_student(
    roll_no: str = Form(...),
    name: str = Form(...),
    email: str = Form(...),
    phone: str = Form(...),
    department: str = Form(...),
    academic_year: str = Form(...)
):
    cursor.execute("""
        UPDATE students
        SET student_name=?, email=?, phone=?, department=?, academic_year=?
        WHERE roll_no=?
    """, (name, email, phone, department, academic_year, roll_no))
    conn.commit()

    json_path = os.path.join(DETAILS_DIR, f"{roll_no}.json")
    if os.path.exists(json_path):
        with open(json_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        data.update({
            "name": name,
            "email": email,
            "phone": phone,
            "department": department,
            "academic_year": academic_year
        })

        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4)

    return {"status": "success", "message": "Student updated successfully"}



# =========================================================
# POST: Delete Student (Complete Cleanup)
# =========================================================
import pandas as pd

from openpyxl import load_workbook
import calendar

def delete_student_from_attendance_excel(roll_no: str):
    if not os.path.exists(ATTENDANCE_DIR):
        return

    for file in os.listdir(ATTENDANCE_DIR):
        if not file.startswith("Attendance-") or not file.endswith(".xlsx"):
            continue

        path = os.path.join(ATTENDANCE_DIR, file)
        wb = load_workbook(path)
        ws = wb.active

        year = int(file.replace("Attendance-", "").replace(".xlsx", ""))

        for month in range(1, 13):
            month_header = f"{calendar.month_name[month].upper()} {year}"

            header_row = None
            for r in range(1, ws.max_row + 1):
                if ws.cell(r, 1).value == month_header:
                    header_row = r
                    break

            if not header_row:
                continue

            r = header_row + 2  # start of students
            while ws.cell(r, 1).value:
                if str(ws.cell(r, 1).value).strip() == roll_no:
                    ws.delete_rows(r, 1)
                    break
                r += 1

        wb.save(path)


@router.post("/delete-student")
def delete_student(
    roll_no: str = Form(...)):
    # DB
    cursor.execute("DELETE FROM students WHERE roll_no=?", (roll_no,))
    conn.commit()

    # Excel
    delete_student_from_attendance_excel(roll_no)

    
    # JSON
    json_file = os.path.join(DETAILS_DIR, f"{roll_no}.json")
    if os.path.exists(json_file):
        os.remove(json_file)

    # Images
    img_dir = os.path.join(IMAGES_DIR, roll_no)
    if os.path.exists(img_dir):
        shutil.rmtree(img_dir)

    # Encoding
    enc_file = os.path.join(ENCODINGS_DIR, f"{roll_no}.pkl")
    if os.path.exists(enc_file):
        os.remove(enc_file)

    # DOCX
    docx_file = os.path.join(DOCUMENTS_DIR, f"{roll_no}.docx")
    if os.path.exists(docx_file):
        os.remove(docx_file)

    # Reload face encodings
    reload_encodings(known_encodings, known_rolls)

    return {
        "status": "success",
        "message": f"Student '{roll_no}' deleted completely"
    }



from datetime import datetime
from fastapi import APIRouter

cursor.execute("""
CREATE TABLE IF NOT EXISTS academic_year_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    previous_year TEXT NOT NULL,
    new_year TEXT NOT NULL,
    promoted_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
)
""")
conn.commit()

@router.post("/students/promote-academic-year")
def promote_academic_year():
    cursor.execute("SELECT DISTINCT academic_year FROM students")
    rows = cursor.fetchall()

    if not rows:
        return {"status": "error", "message": "No students found"}

    promoted_pairs = set()
    updated = 0

    for (year_str,) in rows:
        if not year_str or "-" not in year_str:
            continue

        start, end = year_str.split("-")
        start, end = int(start), int(end)

        new_year = f"{start + 1}-{str(end + 1).zfill(2)}"
        promoted_pairs.add((year_str, new_year))

        cursor.execute("""
            SELECT COUNT(*) FROM academic_year_history
            WHERE DATE(promoted_at) = DATE('now')
        """)
        if cursor.fetchone()[0] > 0:
            return {
                "status": "error",
                "message": "Academic year already promoted today"
            }

        cursor.execute(
            "UPDATE students SET academic_year=? WHERE academic_year=?",
            (new_year, year_str)
        )
        updated += cursor.rowcount

    # Save history (only once)
    for prev, new in promoted_pairs:
        cursor.execute(
            "INSERT INTO academic_year_history (previous_year, new_year) VALUES (?, ?)",
            (prev, new)
        )

    conn.commit()

    return {
        "status": "success",
        "message": f"Academic year promoted for {updated} students",
        "new_year": new_year
    }

@router.post("/students/undo-academic-year")
def undo_academic_year():
    cursor.execute("""
        SELECT previous_year, new_year
        FROM academic_year_history
        ORDER BY promoted_at DESC
        LIMIT 1
    """)
    row = cursor.fetchone()

    if not row:
        return {
            "status": "error",
            "message": "No promotion history found to undo"
        }

    prev_year, new_year = row

    # Revert students
    cursor.execute(
        "UPDATE students SET academic_year=? WHERE academic_year=?",
        (prev_year, new_year)
    )
    reverted = cursor.rowcount

    # Remove this history entry (single-level undo)
    cursor.execute("""
        DELETE FROM academic_year_history
        WHERE previous_year=? AND new_year=?
    """, (prev_year, new_year))

    conn.commit()

    return {
        "status": "success",
        "message": f"Undo successful. Reverted {reverted} students to {prev_year}",
        "restored_year": prev_year
    }
