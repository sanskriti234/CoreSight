# =========================================================
#               CoreSight Attendance Module
# =========================================================

import os
import cv2
import sqlite3
import calendar
import numpy as np
import face_recognition

from datetime import datetime
from fastapi import APIRouter, UploadFile, File

from openpyxl import Workbook, load_workbook
from paths import ATTENDANCE_DIR
from face_service import known_encodings, known_rolls

# =========================================================
# Router
# =========================================================

router = APIRouter(prefix="/attendance", tags=["Attendance"])

# =========================================================
# Database
# =========================================================

DB_PATH = "Core.db"

conn = sqlite3.connect(DB_PATH, check_same_thread=False)
cursor = conn.cursor()

cursor.execute(""" CREATE TABLE IF NOT EXISTS attendance_logs (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    roll_no TEXT,
    date TEXT,
    time TEXT,
    status TEXT
)
""")

conn.commit()

# =========================================================
# Create Attendance Excel
# =========================================================

def create_attendance_excel(year):

    file_path = os.path.join(
        ATTENDANCE_DIR,
        f"Attendance-{year}.xlsx"
    )

    if os.path.exists(file_path):
        return file_path
    wb = Workbook()
    ws = wb.active

    row = 1

    for month in range(1, 13):

        month_name = f"{calendar.month_name[month].upper()} {year}"

        ws.cell(row=row, column=1).value = month_name
        row += 1

         # Headers
        ws.cell(row=row, column=1).value = "Roll No"

        days = calendar.monthrange(year, month)[1]

        for d in range(1, days + 1):
            ws.cell(row=row, column=d + 1).value = d

        row += 1

        # Reserve rows for students
        row += 40

    wb.save(file_path)

    return file_path

# =========================================================
# Mark Attendance in Excel
# =========================================================

def mark_attendance_excel(roll_no):

    today = datetime.now()

    year = today.year
    month = today.month
    day = today.day

    file_path = create_attendance_excel(year)

    wb = load_workbook(file_path)
    ws = wb.active

    month_header = f"{calendar.month_name[month].upper()} {year}"

    header_row = None

    for r in range(1, ws.max_row + 1):

        if ws.cell(r, 1).value == month_header:
            header_row = r
            break

    if not header_row:
        return False

    student_row = None

    r = header_row + 2

    while ws.cell(r, 1).value:

        if str(ws.cell(r, 1).value).strip() == roll_no:
            student_row = r
            break

        r += 1

    # Add new student automatically
    if not student_row:
        student_row = r
        ws.cell(student_row, 1).value = roll_no

    attendance_col = day + 1

    # Prevent duplicate attendance
    if ws.cell(student_row, attendance_col).value == "P":
        return "already_marked"

    ws.cell(student_row, attendance_col).value = "P"

    wb.save(file_path)

    return "marked"

# =========================================================
# Face Recognition Attendance API
# =========================================================

@router.post("/mark-by-face")
async def mark_by_face(frame: UploadFile = File(...)):

    image_bytes = await frame.read()

    np_arr = np.frombuffer(image_bytes, np.uint8)

    image = cv2.imdecode(np_arr, cv2.IMREAD_COLOR)

    rgb = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)

    face_locations = face_recognition.face_locations(rgb)

    if not face_locations:
        return {"status": "no_face"}

    face_encodings = face_recognition.face_encodings(
        rgb,
        face_locations
    )

    for enc in face_encodings:

        distances = face_recognition.face_distance(
            known_encodings,
            enc
        )

        if len(distances) == 0:
            continue

        best_match = np.argmin(distances)

        if distances[best_match] < 0.45:

            roll_no = known_rolls[best_match]

            attendance_result = mark_attendance_excel(roll_no)

            now = datetime.now()

            if attendance_result == "already_marked":
                return {
                    "status": "already_marked",
                    "roll_no": roll_no
                }

            cursor.execute("""
            INSERT INTO attendance_logs
            (roll_no, date, time, status)
            VALUES (?, ?, ?, ?)
            """, (
                roll_no,
                now.strftime("%Y-%m-%d"),
                now.strftime("%H:%M:%S"),
                "Present"
            ))

            conn.commit()

            return {
                "status": "marked_present",
                "roll_no": roll_no
            }
        
    return {"status": "unknown"}
# =========================================================
# Today's Logs API
# =========================================================


@router.get("/logs-today")
def logs_today():

    today = datetime.now().strftime("%Y-%m-%d")

    cursor.execute("""
    SELECT roll_no, date, time, status
    FROM attendance_logs
    WHERE date=?
    ORDER BY id DESC
    """, (today,))

    rows = cursor.fetchall()

    logs = []

    for row in rows:
        logs.append({
            "roll_no": row[0],
            "date": row[1],
            "time": row[2],
            "status": row[3]
        })

    return {"logs": logs}